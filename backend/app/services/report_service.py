"""PDF and Excel report generation from a stored analysis."""

import re
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.athlete import Athlete
from app.models.video import Analysis

REPORT_TITLES = {
    "injury_risk": "Injury Risk Report",
    "biomechanical": "Biomechanical Assessment Report",
    "movement": "Movement Analysis Report",
    "performance": "Athlete Performance Report",
    "rehabilitation": "Rehabilitation Report",
}


def _fmt(v) -> str:
    if v is None:
        return "-"
    if isinstance(v, float):
        return f"{v:.2f}".rstrip("0").rstrip(".")
    return str(v)


def _sections(report_type: str, athlete: Athlete, analysis: Analysis, history: list[tuple], injuries: list[dict]) -> list[dict]:
    """Report content as neutral structures shared by the PDF and Excel writers."""
    bio, scores = analysis.biomechanics, analysis.scores
    out: list[dict] = [
        {
            "title": "Summary",
            "table": [
                ["Athlete", athlete.full_name], ["Athlete ID", athlete.athlete_code],
                ["Sport / position", f"{athlete.sport_type} / {athlete.position or '-'}"],
                ["Activity analysed", analysis.activity],
                ["Analysis date", analysis.created_at.strftime("%Y-%m-%d %H:%M")],
                ["Injury risk score", f"{analysis.risk_score} ({analysis.risk_category})"],
                ["Movement quality", _fmt(analysis.movement_quality)],
            ],
        }
    ]

    if report_type == "injury_risk":
        out.append({
            "title": "Risk by injury category",
            "header": ["Category", "Score", "Level", "Main drivers"],
            "table": [[r["category"], r["score"], r["level"], ", ".join(d["driver"] for d in r["drivers"][:3]) or "-"]
                      for r in analysis.injury_risks],
        })
        risk = scores.get("risk", {})
        out.append({
            "title": "Score components (weighted)",
            "header": ["Component", "Score (0-100)", "Weight used"],
            "table": [[k.replace("_", " "), v, risk.get("weights_used", {}).get(k)] for k, v in risk.get("components", {}).items()],
        })
        prob = scores.get("injury_probability", {})
        out.append({"title": "Estimated injury probability",
                    "paragraphs": [f"{_fmt(prob.get('probability'))} (source: {prob.get('source', '-')}). "
                                   "This is a screening estimate, not a clinical prediction."]})

    elif report_type == "biomechanical":
        out.append({"title": "Biomechanical metrics", "header": ["Metric", "Value"],
                    "table": [[k.replace("_", " "), v] for k, v in bio["metrics"].items()]})
        out.append({
            "title": "Joint angles (degrees)",
            "header": ["Joint", "Left mean", "Left ROM", "Right mean", "Right ROM"],
            "table": [[k.replace("_", " "), v["left"]["mean"], v["left"]["rom"], v["right"]["mean"], v["right"]["rom"]]
                      for k, v in bio["joint_angles"].items()],
        })
        out.append({"title": "Left-right symmetry index (%)", "header": ["Measure", "Value"],
                    "table": [[k.replace("_", " "), v] for k, v in bio["symmetry"].items()]})
        if bio.get("landing"):
            out.append({"title": "Landing mechanics", "header": ["Measure", "Value"],
                        "table": [[k.replace("_", " "), v] for k, v in bio["landing"].items()]})

    elif report_type == "movement":
        out.append({
            "title": "Movement findings",
            "header": ["Measure", "Value", "Level", "Note"],
            "table": [[d["label"], d["value"], d["level"], d["message"] or ""] for d in analysis.deviations],
        })
        an = analysis.anomalies
        rows = []
        for key, label in (("inconsistency", "Rep-to-rep inconsistency"), ("fatigue", "Fatigue drift"), ("performance_decline", "Performance decline")):
            data = an.get(key)
            rows.append([label, "; ".join(f"{k}: {_fmt(v)}" for k, v in data.items()) if data else "Not enough data"])
        out.append({"title": "Anomaly detection", "header": ["Check", "Result"], "table": rows})

    elif report_type == "performance":
        out.append({
            "title": "Scores",
            "header": ["Score", "Value"],
            "table": [[k.replace("_", " "), scores.get(k)] for k in ("movement_quality", "biomechanical_efficiency", "fatigue_risk", "overall_health")],
        })
        if history:
            out.append({"title": "Trend (recent analyses)", "header": ["Date", "Risk score", "Movement quality"],
                        "table": [[d.strftime("%Y-%m-%d"), r, q] for d, r, q in history]})

    elif report_type == "rehabilitation":
        recs = analysis.recommendations
        active = [i for i in injuries if i["status"] in ("active", "recovering")]
        out.append({"title": "Active and recovering injuries", "header": ["Injury", "Body part", "Date", "Status"],
                    "table": [[i["injury_type"], i["body_part"], str(i["injury_date"]), i["status"]] for i in active]
                    or [["None recorded", "", "", ""]]})
        rows = [[c["title"], e["name"], e["prescription"], e["purpose"]]
                for c in recs.get("corrective_exercises", []) for e in c["exercises"]]
        out.append({"title": "Corrective exercises", "header": ["Focus", "Exercise", "Prescription", "Purpose"],
                    "table": rows or [["No corrective work flagged", "", "", ""]]})
        out.append({"title": "Training modifications", "paragraphs": recs.get("training_modifications", [])})
        out.append({"title": "Recovery plan", "paragraphs": recs.get("recovery_plan", [])})
        out.append({"title": "Note", "paragraphs": [recs.get("disclaimer", "")]})
    return out


def build_pdf(path: Path, report_type: str, athlete: Athlete, analysis: Analysis, history: list[tuple], injuries: list[dict]) -> None:
    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell")
    cell.fontSize, cell.leading = 8, 10
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm, title=REPORT_TITLES[report_type])
    story = [Paragraph(REPORT_TITLES[report_type], styles["Title"]),
             Paragraph(f"Generated {date.today().isoformat()}", styles["Normal"]), Spacer(1, 6 * mm)]
    for section in _sections(report_type, athlete, analysis, history, injuries):
        story.append(Paragraph(section["title"], styles["Heading2"]))
        for text in section.get("paragraphs", []):
            story.append(Paragraph(escape(str(text)), styles["BodyText"]))
        if "table" in section:
            data = [[Paragraph(escape(_fmt(c)), cell) for c in row] for row in section["table"]]
            if "header" in section:
                data.insert(0, [Paragraph(f"<b>{escape(h)}</b>", cell) for h in section["header"]])
            table = Table(data, repeatRows=1 if "header" in section else 0, hAlign="LEFT")
            style = [("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey), ("VALIGN", (0, 0), (-1, -1), "TOP")]
            if "header" in section:
                style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF7")))
            table.setStyle(TableStyle(style))
            story.append(table)
        story.append(Spacer(1, 4 * mm))
    doc.build(story)


def build_excel(path: Path, report_type: str, athlete: Athlete, analysis: Analysis, history: list[tuple], injuries: list[dict]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="E8EEF7")
    for section in _sections(report_type, athlete, analysis, history, injuries):
        ws = wb.create_sheet(re.sub(r"[\[\]\*\?/\\:]", "", section["title"])[:31])
        row = 1
        for text in section.get("paragraphs", []):
            ws.cell(row=row, column=1, value=str(text)).alignment = Alignment(wrap_text=True)
            row += 1
        if "header" in section:
            for c, h in enumerate(section["header"], start=1):
                cell = ws.cell(row=row, column=c, value=h)
                cell.font, cell.fill = Font(bold=True), header_fill
            row += 1
        for data_row in section.get("table", []):
            for c, v in enumerate(data_row, start=1):
                ws.cell(row=row, column=c, value=v)
            row += 1
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 60)
    wb.save(path)
